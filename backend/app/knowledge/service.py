from io import BytesIO
from pathlib import Path
from uuid import uuid4
import math
import zipfile
from fastapi import HTTPException
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document
from app.config import settings
from app.database import database, audit

MAX_BYTES = 8 * 1024 * 1024


def extract_text(filename: str, content: bytes) -> str:
    if not content or len(content) > MAX_BYTES:
        raise HTTPException(422, 'File must be between 1 byte and 8 MB')
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == '.pdf' and content.startswith(b'%PDF-'):
            pdf = PdfReader(BytesIO(content))
            if pdf.is_encrypted or len(pdf.pages) > 150:
                raise ValueError('Unsupported PDF')
            text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
        elif suffix == '.docx' and content.startswith(b'PK'):
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 24 * 1024 * 1024 or len(archive.infolist()) > 1000:
                    raise ValueError('Oversized archive')
                if any(n.endswith('vbaProject.bin') for n in archive.namelist()):
                    raise ValueError('Macros not permitted')
            doc = Document(BytesIO(content))
            text = '\n'.join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
        elif suffix == '.txt':
            text = content.decode('utf-8-sig')
        else:
            raise ValueError('Unsupported format')
    except Exception:
        raise HTTPException(422, 'Use a readable, unencrypted PDF, DOCX, or UTF-8 TXT file') from None
    text = text.replace('\x00', '').strip()
    if not text or len(text) > 250000:
        raise HTTPException(422, 'No extractable text, or document exceeds 250,000 characters. Scanned PDFs require OCR first.')
    return text


def chunk_text(text, size=1800, overlap=250):
    return [text[i:i+size] for i in range(0, len(text), size-overlap) if text[i:i+size].strip()]


def gemini():
    return genai.Client(api_key=settings().gemini_api_key.get_secret_value(),
                        http_options=types.HttpOptions(timeout=60000))


def embed(texts, task='RETRIEVAL_DOCUMENT'):
    with gemini() as client:
        result = client.models.embed_content(model=settings().embedding_model, contents=texts,
            config=types.EmbedContentConfig(output_dimensionality=1536, task_type=task))
    vectors = []
    for item in result.embeddings:
        values = item.values
        if len(values) != 1536 or not all(math.isfinite(v) for v in values):
            raise ValueError('Unexpected embedding shape')
        norm = math.sqrt(sum(v*v for v in values))
        vectors.append([v/norm for v in values])
    return vectors


def search_knowledge(query, specialty=None, doctor=None, language=None):
    vector = embed([query], 'RETRIEVAL_QUERY')[0]
    rows = database().rpc('clinic_match_knowledge', {'p_embedding': vector,
        'p_model': settings().embedding_model, 'p_specialty': specialty, 'p_doctor': doctor, 'p_language': language}).execute().data
    return [r for r in rows if r['similarity'] >= 0.35]


def ingest(actor, filename, content, title, doctor_id=None, specialty_id=None, language='ar', document_type='general'):
    text = extract_text(filename, content)
    doc_id = str(uuid4())
    path = f'{actor["id"]}/{doc_id}{Path(filename).suffix.lower()}'
    metadata = {'doctor_id': doctor_id, 'specialty_id': specialty_id,
                'document_type': document_type, 'language': language}
    storage = database().storage.from_('knowledge-base')
    mime = {'.pdf':'application/pdf','.docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','.txt':'text/plain'}[Path(filename).suffix.lower()]
    storage.upload(path, content, file_options={'content-type': mime, 'upsert':'false'})
    try:
        database().table('knowledge_documents').insert({'id':doc_id,'title':title,'file_path':path,
            **metadata,'status':'processing','active':False,'embedding_model':settings().embedding_model}).execute()
    except Exception:
        storage.remove([path])
        raise
    try:
        chunks = chunk_text(text)
        for offset in range(0, len(chunks), 20):
            batch = chunks[offset:offset+20]
            vectors = embed(batch)
            database().table('knowledge_chunks').insert([
                {'document_id':doc_id,'content':value,'embedding':vector,
                 'metadata':{**metadata,'chunk_index':offset+i,'embedding_model':settings().embedding_model}}
                for i,(value,vector) in enumerate(zip(batch,vectors,strict=True))]).execute()
        row = database().table('knowledge_documents').update({'status':'completed','active':True}).eq('id',doc_id).execute().data[0]
        audit(actor['id'],'ingest','knowledge_documents',doc_id)
        return row
    except Exception:
        database().table('knowledge_documents').update({'status':'failed','active':False}).eq('id',doc_id).execute()
        raise HTTPException(503, 'Document saved but indexing failed. Retry or replace it from Knowledge.') from None
